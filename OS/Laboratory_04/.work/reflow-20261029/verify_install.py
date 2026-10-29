from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
import json,hashlib,shutil,re
W=Path(__file__).resolve().parent;R=W.parents[1]
original=R/'Отчёт_ЛР4.docx';backup=W/'Отчёт_ЛР4 — до исправления переносов.docx';ready=W/'Отчёт_ЛР4 — исправленный.docx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(original)==sha(backup),'Source changed during review'
old=Document(backup);d=Document(ready)
assert len(d.inline_shapes)==25 and len(d.tables)==12
def media(p):
 with ZipFile(p) as z:return Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
assert media(backup)==media(ready)
def values(t):return [[c.text for c in row.cells] for row in t.rows]
for a,b in zip(old.tables,d.tables):assert values(a)==values(b)
assert [''.join(e.itertext()) for e in old.element.xpath('.//m:oMath')]==[''.join(e.itertext()) for e in d.element.xpath('.//m:oMath')]
def effective(p,name):
 value=getattr(p.paragraph_format,name);style=p.style
 while value is None and style is not None:value=getattr(style.paragraph_format,name);style=style.base_style
 return value
seen=False;figs=0;bodybreaks=0
for p in d.paragraphs:
 if p.text=='Введение' and p.style.name=='Report unnumbered':seen=True
 if not seen:continue
 chapter=p.style.name in ['Heading 1','Report unnumbered']
 assert not p._p.xpath('.//w:br[@w:type="page"]'),p.text
 if effective(p,'page_break_before'):
  bodybreaks+=1;assert chapter,p.text
 if p._p.xpath('.//w:drawing'):
  figs+=1;assert effective(p,'space_before').pt==14
  cap=Paragraph(p._p.getnext(),d._body);after=Paragraph(cap._p.getnext(),d._body)
  assert cap.text.startswith('Рисунок ') and effective(cap,'keep_with_next')
  assert after.text.strip() and not after.text.startswith(('Рисунок ','Таблица ','Листинг '))
  assert not effective(after,'keep_together')
assert figs==25
tables=json.loads((W/'word-tables.json').read_text(encoding='utf-8-sig'))
assert len(tables)==12 and all(t['FirstPage']==t['LastPage'] for t in tables)
audit=json.loads((W/'page-audit.json').read_text(encoding='utf8'))
assert not audit['issues'] and len(audit['pages'])==43
assert len([t for t in d.tables if len(t.rows)==len(t.columns)==1])==6
oldwords=len(' '.join(p.text for p in old.paragraphs).split());newwords=len(' '.join(p.text for p in d.paragraphs).split())
assert newwords-oldwords>600
report={'pages':43,'added_body_words':newwords-oldwords,'screenshots_unchanged':True,'all_table_contents_unchanged':True,'native_equation_preserved':True,'all_pages_visually_reviewed':True,'manual_body_page_breaks':0,'forced_chapter_starts':bodybreaks,'layout_issues':audit['issues'],'sha256':sha(ready),'output':str(original),'backup':str(backup)}
shutil.copy2(ready,original)
assert sha(original)==sha(ready)
(W/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
