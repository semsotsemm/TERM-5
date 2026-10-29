from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import hashlib,json,re,shutil
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import pypdfium2 as pdfium
from PIL import ImageChops

W=Path(__file__).resolve().parent;R=W.parents[1]
original=R/'Отчёт_ЛР4.docx';backup=W/'Отчёт_ЛР4 — до оформления.docx';ready=R/'Отчёт_ЛР4 — оформление.docx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(original)==sha(backup),'Source was edited during review; do not overwrite.'
d=Document(ready);old=Document(backup)
assert len(d.inline_shapes)==25 and len(d.tables)==12
def media(p):
 with ZipFile(p) as z:return Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
assert media(backup)==media(ready)
assert not d.element.xpath('.//wp:anchor')
def level(p):
 if p.style.name in ['Heading 1','Report unnumbered']:return 1
 if p.style.name in ['Heading 2','Report appendix subsection']:return 2
 if p.style.name=='Heading 3':return 3
 return 0
def effective(p,name):
 value=getattr(p.paragraph_format,name)
 style=p.style
 while value is None and style is not None:
  value=getattr(style.paragraph_format,name);style=style.base_style
 return value
caps=[]
for p in d.paragraphs:
 if level(p):
  br=effective(p,'page_break_before')
  assert bool(br)==(level(p)==1),(p.text,br)
  e=p._p.getprevious()
  if e is not None and e.tag==qn('w:p'):
   assert not Paragraph(e,d._body).text.startswith(('Рисунок ','Таблица ','Листинг '))
  nxt=p._p.getnext()
  if nxt is not None and nxt.tag==qn('w:p'):
   np=Paragraph(nxt,d._body)
   expected=(12 if level(p)==1 else 6) if level(np) else (18 if level(p)==1 else 12)
   assert effective(p,'space_after').pt==expected,(p.text,expected)
 if p._p.xpath('.//w:drawing'):
  assert effective(p,'space_before').pt==14 and effective(p,'keep_with_next')
  cp=Paragraph(p._p.getnext(),d._body);after=Paragraph(cp._p.getnext(),d._body)
  assert cp.text.startswith('Рисунок ') and effective(cp,'keep_with_next')
  assert after.text.strip() and not after.text.startswith(('Рисунок ','Таблица ','Листинг ')) and not level(after)
  caps.append(cp.text)
assert len(caps)==25 and len(set(caps))==25
def values(t):return [[c.text for c in r.cells] for r in t.rows]
for i in [0,1,3,4,5]:assert values(old.tables[i])==values(d.tables[i]),i
for i in range(6,12):
 t=d.tables[i]
 assert len(t.rows)==len(t.columns)==1
 assert re.sub(r'\s+','',old.tables[i].cell(0,0).text)==re.sub(r'\s+','',t.cell(0,0).text)
 for side in ['top','left','bottom','right']:
  margin=t._tbl.tblPr.find(qn('w:tblCellMar')).find(qn('w:'+side))
  assert 169<=int(margin.get(qn('w:w')))<=171
 assert t.rows[0]._tr.trPr.find(qn('w:cantSplit')) is not None
assert [''.join(e.itertext()) for e in old.element.xpath('.//m:oMath')]==[''.join(e.itertext()) for e in d.element.xpath('.//m:oMath')]
audit=json.loads((W/'page-audit.json').read_text(encoding='utf8'))
assert not audit['issues'] and len(audit['pages'])==41
tables=json.loads((W/'word-tables.json').read_text(encoding='utf-8-sig'))
assert len(tables)==12 and all(t['FirstPage']==t['LastPage'] for t in tables)
before=pdfium.PdfDocument(W/'before.pdf');after=pdfium.PdfDocument(W/'final.pdf')
a=before[0].render(scale=1.7).to_pil().convert('RGB');b=after[0].render(scale=1.7).to_pil().convert('RGB')
assert ImageChops.difference(a,b).getbbox() is None
report={'pages':41,'screenshots':25,'listing_tables_1x1':6,'data_tables':5,'screenshots_unchanged':True,'code_tokens_unchanged':True,'data_tables_unchanged':True,'cover_unchanged':True,'all_pages_visually_reviewed':True,'layout_issues':audit['issues'],'output':str(original),'backup':str(backup),'sha256':sha(ready)}
archive=W/'final-reviewed.docx'
assert not archive.exists()
shutil.copy2(ready,original)
assert sha(original)==report['sha256']
ready.replace(archive)
(W/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
