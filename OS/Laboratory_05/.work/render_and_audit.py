from pathlib import Path
import json
import re
import pypdfium2 as pdfium
import pdfplumber
from PIL import Image,ImageDraw
from docx import Document

root=Path(__file__).resolve().parent
pdf=pdfium.PdfDocument(root/'report.pdf')
pages=root/'report-pages'; pages.mkdir(exist_ok=True)
for i,page in enumerate(pdf):
    page.render(scale=1.5).to_pil().save(pages/f'page-{i+1:02}.png')
for start in range(0,len(pdf),6):
    canvas=Image.new('RGB',(1500,3*730),'#bbb'); draw=ImageDraw.Draw(canvas)
    for j in range(start,min(start+6,len(pdf))):
        im=Image.open(pages/f'page-{j+1:02}.png'); im.thumbnail((490,690))
        x=((j-start)%2)*750; y=((j-start)//2)*730
        canvas.paste(im,(x+90,y+25)); draw.text((x+5,y+4),str(j+1),fill='black')
    canvas.save(root/f'report-sheet-{start//6+1}.jpg')
doc=pdfplumber.open(root/'report.pdf').pages
inventory=[]
for i,page in enumerate(doc):
    text=page.extract_text() or ''
    lines=[l for l in text.splitlines() if l.strip()]
    figs=re.findall(r'Рисунок (\d+\.\d+)',text)
    markers=re.findall(r'СКРИНШОТА (\d+\.\d+)',text)
    inventory.append(dict(page=i+1,lines=len(lines),start=lines[:3],end=lines[-3:],figures=figs,markers=markers))
    for f in figs:
        assert f in markers, ('Separated figure marker/caption',i+1,f)
(root/'page-inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
(root/'report-text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n'+(p.extract_text() or '') for i,p in enumerate(doc)),encoding='utf-8')
d=Document(root.parent/'Отчет_ЛР5.docx')
assert len(re.findall(r'\[МЕСТО ДЛЯ СКРИНШОТА', '\n'.join(p.text for p in d.paragraphs)))==37
assert len(d.tables)==9
print('Pages',len(pdf),'figures',sum(len(p['figures']) for p in inventory))
for p in inventory: print(p['page'],str(p['start'])[:120],p['figures'])
