from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from PIL import Image,ImageOps,ImageDraw
import json
R=Path(__file__).resolve().parents[1]
W=R/'.work/review-20261027';W.mkdir(exist_ok=True)
p=R/'Отчет_ЛР5.docx';d=Document(p)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
items=[];images=[]
for i,e in enumerate(d.element.body):
 t=''.join(e.xpath('.//w:t/text()'))
 if e.tag.endswith('}tbl'):
  t='TABLE '+json.dumps([[''.join(c.xpath('.//w:t/text()')) for c in row.xpath('./w:tc')] for row in e.xpath('./w:tr')],ensure_ascii=False)
 im=[]
 for b in e.xpath('.//a:blip'):
  rid=b.get('{'+ns['r']+'}embed')
  if not rid:continue
  part=d.part.related_parts[rid]; dest=W/(f'img{len(images)+1:02d}'+Path(part.partname).suffix);dest.write_bytes(part.blob)
  obj={'index':len(images)+1,'body':i,'rid':rid,'path':str(dest),'size':Image.open(dest).size}
  images.append(obj);im.append(obj['index'])
 hi=e.xpath('.//w:highlight/@w:val')
 items.append({'body':i,'text':t,'style':e.xpath('./w:pPr/w:pStyle/@w:val'),'images':im,'highlights':hi})
(W/'inventory.json').write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf8')
(W/'images.json').write_text(json.dumps(images,ensure_ascii=False,indent=2),encoding='utf8')
(W/'text.txt').write_text('\n'.join(f"[{x['body']}] {'IMAGE'+str(x['images']) if x['images'] else ''} {'HIGHLIGHT' if x['highlights'] else ''} {' '.join(x['style'])} {x['text']}" for x in items),encoding='utf8')
sd=Document(R/'ОС ЛР5.docx');(W/'assignment.txt').write_text('\n'.join(p.text for p in sd.paragraphs)+'\n'+'\n'.join(str([[c.text for c in r.cells] for r in t.rows]) for t in sd.tables),encoding='utf8')
for start in range(0,len(images),6):
 batch=images[start:start+6];sheet=Image.new('RGB',(1800,3*480),'#ddd');dr=ImageDraw.Draw(sheet)
 for j,x in enumerate(batch):
  im=Image.open(x['path']).convert('RGB');im.thumbnail((880,440));xx=(j%2)*900;yy=(j//2)*480;sheet.paste(im,(xx,yy+30));dr.text((xx+10,yy+8),f"IMAGE {x['index']:02d} BODY {x['body']} {x['size']}",fill='black')
 sheet.save(W/f"sheet-{start//6+1:02d}.jpg")
print(json.dumps({'body':len(items),'images':len(images),'tables':len(d.tables),'work':str(W)},ensure_ascii=False))
