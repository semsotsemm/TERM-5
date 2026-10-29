from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import ast,hashlib,json,re,shutil
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

R=Path(__file__).resolve().parents[2];W=Path(__file__).resolve().parent
SRC=R/'Отчет_ЛР5.docx';OUT=R/'Отчет_ЛР5 — готовый.docx'
BACK=W/'Отчет_ЛР5 — до доработки.docx'
if not BACK.exists():shutil.copy2(SRC,BACK)
d=Document(BACK);body=list(d.element.body)
P=lambda i:Paragraph(body[i],d._body)
def setp(i,t):
 p=P(i);p.text=t;return p
def remove(i):
 e=body[i]
 if e.getparent() is not None:e.getparent().remove(e)

repl={
67:'Цель работы — получить практические навыки управления планированием процессов и потоков в Windows и Linux и исследовать алгоритмы распределения процессорного времени на заданных наборах задач.',
69:'Для достижения цели выполнены следующие задачи:',
70:'– разработаны приложения Lab-05a, Lab-05x, Lab-05b и Lab-05c для Windows и Linux;',
71:'– проведены опыты на всех доступных процессорах и на одном процессоре;',
72:'– реализованы SJF, Round Robin и вытесняющее приоритетное планирование;',
73:'– сопоставлена статистика при различных квантах времени;',
74:'– реализована многоуровневая очередь с обратной связью для 100 процессов.',
75:'Натурные опыты выполнены в Windows и Linux с одинаковой периодической задержкой. Для SJF, RR, CP и MQ построены модельные расписания и рассчитаны характеристики обслуживания без выполнения реальных пользовательских заданий.',
87:'Исследование выполнено в гостевых Windows 10 Pro x64 версии 10.0.19045 и Ubuntu 26.04.1 LTS с ядром 7.0.0-31-generic. В обеих средах отображается AMD Ryzen 7 7735H with Radeon Graphics; доступность восьми логических CPU подтверждена масками приложений. Для Linux разрешены CPU 0–7.',
90:'В Windows получена версия 10.0.19045 и 64-разрядная архитектура. Запрос Win32_Processor показывает восемь ядер; доступность восьми логических CPU для исследуемого процесса дополнительно подтверждается Lab-05a.',
93:'В Ubuntu разрешённый набор состоит из CPU 0–7. Полная маска равна 255, то есть 11111111₂, маска одного процессора — 1. В опытах с маской 1 используется CPU 0. Снимок показывает один аппаратный поток на ядро в конфигурации гостевой системы.',
101:'В контрольном запуске PID равен 2484, TID — 5660. Получены класс Normal (32), THREAD_REL = 0 и BASE = 8. Маски процесса и системы равны 255, CPU_COUNT = 8, текущий CPU = 1. GetThreadPriority возвращает относительный уровень; поле BASE содержит вычисленный базовый приоритет.',
106:'Процесс PID = 6116, TID = 7076 выполнил 1 000 000 итераций. WALL_SECONDS = 204,716943 с, CLOCK_SECONDS = 204,717000 с. Близость значений соответствует реализации clock в Microsoft CRT. В последних строках встречаются CPU 1, 2 и 3, что подтверждает миграции в пределах разрешённой маски.',
114:'Для текущего запуска родитель имеет PID = 9484, дочерние процессы — PID = 12272 и 9868. В Process Explorer оба дочерних процесса имеют приоритет 8. Их классы Normal и относительные уровни 0 соответствуют одинаковому базовому приоритету основных потоков.',
117:'Первым зарегистрировано завершение второго процесса: PID = 9868, TID = 1300, SLOT = 2. В этот момент ITER1 = 999000, ITER2 = 1000000, DELTA = 1000. Время второго процесса составило 203,312802 с, первого — 203,637669 с. При одинаковых приоритетах небольшое расхождение связано с конкретным порядком обслуживания и выводом; оно не указывает на устойчивое преимущество одного процесса.',
123:'Родительский PID = 1064, дочерние PID = 7136 и 5180. Для Below Normal показаны CLASS_CODE = 16384 и BASE = 6, для High — CLASS_CODE = 128 и BASE = 13. Значения согласуются с Process Explorer; в выводе видны разные номера разрешённых CPU.',
126:'Первым завершился второй процесс: PID = 5180, TID = 6924, SLOT = 2. Зарегистрированы ITER1 = 1000000, ITER2 = 1000000 и DELTA = 0. Времена High и Below Normal равны 203,677113 и 203,762687 с соответственно. Разные моменты завершения при равных счётчиках допустимы: последняя задержка выполняется после установки счётчика в 1 000 000.',
132:'Родительский PID = 12176, дочерние PID = 6088 и 8088. Маска 1 разрешает только CPU 0, что подтверждается окном Process Affinity и рабочими строками. Основные потоки Below Normal и High имеют базовые приоритеты 6 и 13 соответственно.',
135:'Первым завершился процесс High: PID = 8088, TID = 2036, SLOT = 2. На сохранённом снимке ITER1 = 998000, ITER2 = 1000000 и DELTA = 2000. Время High составило 203,706478 с, Below Normal — 204,066889 с. В данном запуске высокоприоритетный процесс опередил второй на одном CPU; один опыт не определяет постоянную величину выигрыша.',
143:'Оба рабочих потока принадлежат PID = 11756 и имеют TID = 7996 и 5412. В Process Explorer для этих TID показан Base Pri = 8; рабочие строки содержат THREAD_REL = 0 и BASE = 8. Использование разных CPU соответствует полной маске 255.',
146:'В FIRST_FINISH зарегистрирован SLOT = 2: PID = 11756, TID = 5412. Оба счётчика равны 1000000, DELTA = 0. Время потока TID = 5412 составило 203,208293 с, TID = 7996 — 203,208953 с. Равные параметры и периодическое ожидание привели к почти совпавшему прогрессу.',
152:'В PID = 6776 поток TID = 7948 имеет THREAD_REL = −2 и BASE = 6, а TID = 4720 — THREAD_REL = 2 и BASE = 10. Эти базовые значения подтверждены Process Explorer. Для TID = 4720 текущий Dyn Pri = 11 отличается от Base Pri = 10; динамическое повышение не меняет заданный относительный уровень.',
155:'Первым зарегистрировано завершение потока Highest: PID = 6776, TID = 4720, SLOT = 2. Получены ITER1 = 1000000, ITER2 = 1000000, DELTA = 0. Его WALL_SECONDS = 203,167705 с, у Lowest — 203,167837 с. При доступности нескольких CPU изменение приоритета не привело к различию зафиксированных счётчиков.',
161:'Оба рабочих потока PID = 9388 закреплены за CPU 0 маской 1. TID = 6712 имеет базовый приоритет 6, TID = 9652 — 10. Process Explorer и рабочие строки подтверждают заданные Lowest и Highest и ограничение родственности.',
164:'Первым завершился Highest: PID = 9388, TID = 9652, SLOT = 2. Счётчики ITER1 и ITER2 равны 1000000, DELTA = 0. Времена Highest и Lowest составили 203,331642 и 203,331915 с. Конкуренция на одном CPU не создала измеримого расхождения итераций: исполнители большую часть времени ожидали таймер.',
171:'При одинаковых классах процессов получено расхождение 1000 итераций. Для Below Normal / High оно равно 0 на всех CPU и 2000 на одном CPU. Во всех трёх процессных опытах первым зарегистрирован второй исполнитель. Эти значения характеризуют отдельные запуски и не задают коэффициент ускорения.',
175:'Во всех потоковых опытах DELTA = 0, первым зарегистрирован второй рабочий поток. Приоритеты подтверждены независимо в Process Explorer. Переход к одному CPU сохранил равенство счётчиков; результат согласуется с преобладанием ожидания таймера над вычислительной работой.',
177:'В Windows наибольшее наблюдаемое расхождение составило 2000 итераций, то есть 0,2 % заданного объёма работы. Различие процессов и потоков в этой серии не доказывает преимущества одного способа организации: проведено по одному запуску на вариант, а длительность определяется главным образом задержками и условиями гостевой ОС.',
184:'В запуске PID = TID = 7622 получены NICE = 0, маска 255, CPU_COUNT = 8 и CPU = 2. Политика POLICY = 0 соответствует SCHED_OTHER, SCHED_PRIORITY = 0. Для этой политики вес обслуживания определяется nice; отдельный приоритет реального времени не задан.',
189:'Процесс PID = TID = 7644 завершил 1 000 000 итераций: WALL_SECONDS = 200,187958 с, CLOCK_SECONDS = 0,090033 с. В отличие от календарного времени, clock в Linux отражает CPU-время процесса и почти не учитывает ожидание таймера. В показанном фрагменте используется CPU 0; это не доказывает ограничения маски одним CPU.',
196:'Родитель имеет PID = 7820, дочерние процессы — 7821 и 7822. У обоих NI = 0, Cpus_allowed = ff и Cpus_allowed_list = 0–7. В /proc/PID/stat показаны priority = 20 и nice = 0, тогда как ps выводит PRI = 19; эти представления приоритета не следует смешивать.',
199:'Первым зарегистрирован второй процесс: PID = TID = 7822, SLOT = 2. Сохранены ITER1 = 1000000, ITER2 = 1000000, DELTA = 0. При одинаковом nice и восьми разрешённых CPU оба процесса достигли одинакового прогресса к моменту первого завершения.',
205:'Родительский PID = 8066, дочерние PID = 8067 и 8068. Для них подтверждены NI = 19 и −20, в /proc — nice = 19 и −20, priority = 39 и 0 соответственно. Обе маски равны ff и разрешают CPU 0–7; отрицательное nice действительно установлено.',
208:'Первым зарегистрирован процесс с nice = 19: PID = TID = 8067, SLOT = 1. При этом ITER1 = 1000000, ITER2 = 1000000, DELTA = 0. Такой порядок допустим: высокий вес CPU не отменяет блокировку на таймере и не гарантирует более раннего завершения при наличии свободных CPU.',
214:'Родительский PID = 8186, дочерние PID = 8187 и 8188. Значения NI = 19 и −20 совпадают с параметрами запуска. В обеих записях /proc показаны Cpus_allowed = 01 и Cpus_allowed_list = 0; рабочие строки также содержат CPU = 0.',
217:'Первым зарегистрирован процесс с nice = −20: PID = TID = 8188, SLOT = 2. Оба счётчика равны 1000000, DELTA = 0. Ограничение одним CPU не создало отставания по счётчику: периоды ожидания освобождали процессор для другого исполнителя.',
225:'Родительский PID = 8305; рабочие потоки имеют TID = 8306 и 8307. Команда ps -L показывает NI = 0 для обоих, отдельные записи /proc/8305/task/TID подтверждают маску ff и CPU 0–7. Общий Tgid = 8305 отличает потоки одного процесса от отдельных процессов.',
228:'Первым зарегистрирован второй поток: PID = 8305, TID = 8307, SLOT = 2. Значения ITER1 = 1000000 и ITER2 = 1000000 совпали, DELTA = 0. При одинаковом nice существенного различия прогресса в данном запуске не выявлено.',
234:'Рабочие потоки PID = 8388 имеют TID = 8389 и 8390. В ps -L и /proc подтверждены nice = 19 и −20, маска ff и разрешённые CPU 0–7. Главный поток сохраняет nice = 0, поэтому проверка только строки процесса не заменяет проверку отдельных TID.',
237:'Первым зарегистрирован поток с nice = 19: PID = 8388, TID = 8389, SLOT = 1. Сохранены ITER1 = 1000000, ITER2 = 1000000 и DELTA = 0. Различные веса обслуживания не привели к различию счётчиков при нагрузке с регулярными задержками.',
243:'Для PID = 8496 рабочие TID = 8497 и 8498 имеют nice = 19 и −20. Маска 01 и Cpus_allowed_list = 0 подтверждены отдельно для главного и рабочих потоков. В рабочих строках обоих исполнителей показан CPU = 0.',
246:'Первым зарегистрирован поток с nice = −20: PID = 8496, TID = 8498, SLOT = 2. Оба счётчика равны 1000000, DELTA = 0. Совместное использование одного CPU сохранило одинаковый прогресс по принятой метрике.',
253:'Во всех трёх процессных опытах DELTA = 0. При nice 0 / 0 и при маске одного CPU первым зарегистрирован второй процесс; при nice 19 / −20 на всех CPU — первый. Порядок завершения сам по себе не определяет преимущество приоритета при равных счётчиках.',
257:'Во всех трёх потоковых опытах DELTA = 0. Последовательность зарегистрированных первых исполнителей совпала с процессной серией: второй, первый, второй. Проверки отдельных TID подтвердили фактические nice и маски.',
259:'В Linux и процессы, и потоки достигли одинаковых счётчиков при первом завершении во всех вариантах. В Windows у процессов получены DELTA = 1000, 0 и 2000, у потоков — только 0. Эти единичные измерения не позволяют ранжировать ОС по скорости; определения CLOCK_SECONDS различаются, а WALL_SECONDS включает периодическое ожидание.',
299:'Сводный запуск подтвердил выбор кванта 80 для RR и CP и совпадение показателей с таблицами 4.2 и 4.3. Среднее CPU-время равно 62,5, пропускная способность — 0,016 для всех вариантов. Это следствие нулевой цены переключений и отсутствия простоя в принятом наборе задач.',
310:'При seed = 525 завершены все 100 процессов. Модельный интервал составил 28 957 мс, CPU-обслуживание — 28 955 мс, простой — 2 мс. Среднее ожидание равно 17 911,13 мс, среднее оборотное время — 18 519,02 мс. Получены 2047 переключений между разными PID и 397 перемещений между очередями.',
313:'Использование модельного CPU составило 99,9931 %, средний отклик — 1375,28 мс. Показатель использования относится к эмулятору. Баланс времени согласован: 28 955 мс обслуживания и 2 мс простоя дают общий интервал 28 957 мс; оборот процесса включает ожидание готовности, CPU и I/O.',
322:'Каждый файл Q0, Q1 и Q2 содержит 100 записей без заголовка. Для PID = 1 представлены arrival = 121 мс, execution = 179 мс и turnaround = 17 060 мс. CPU-времена по очередям равны 20, 60 и 99 мс; их сумма равна 179 мс. Ожидания 1716, 5217 и 9948 мс дают полное ожидание 16 881 мс. PID = 3 не посещал Q2: visited = 0 и её показатели равны нулю.',
327:'В журнале PID = 55 переведён Q0 → Q1 в момент 22 мс и Q1 → Q2 в момент 160 мс по причине quantum. PID = 51 возвращён Q1 → Q0 после I/O в момент 1544 мс. Для PID = 87 показан переход Q2 → Q1 в момент 1200 мс. Эти записи подтверждают понижение при исчерпании кванта и повышение при возврате.',
328:'Итоговые доли CPU равны 28,61 % для Q0, 24,95 % для Q1 и 46,44 % для Q2. Календарные бюджеты 50/30/20 действуют при наличии готовой работы во всех очередях; пустые очереди передают свой бюджет другим. Поэтому увеличение доли Q2 за весь опыт согласуется с правилами модели. Выбор Q2 после кванта Q2 не увеличивает счётчик перемещений.',
330:'Разработаны и проверены приложения управления приоритетами и родственностью в Windows и Linux. Подтверждены базовые приоритеты Windows 8 / 8 для Normal, 6 / 13 для Below Normal / High и 6 / 10 для Lowest / Highest. В Linux проверены nice 0 / 0 и 19 / −20 по ps и отдельным записям /proc. Полная маска 255 разрешает CPU 0–7; маска 1 ограничивает выполнение CPU 0.',
331:'В Windows при первом завершении процессов получены расхождения 1000, 0 и 2000 итераций, для потоков — 0 во всех вариантах. В Linux DELTA = 0 во всех шести опытах. Периодические задержки по 200 мс определяют основную длительность нагрузки, поэтому изменение приоритета не гарантирует различия счётчиков. Полученные единичные результаты не устанавливают устойчивого преимущества процессов, потоков или одной из ОС.',
332:'Для десяти заданных задач рассчитаны расписания SJF, RR и CP. Среднее ожидание составило 154,0 для SJF, 170,5 для CP и 207,5 для RR при кванте 80. Среди проверенных квантов 80 даёт минимальное ожидание, небольшие кванты улучшают отклик. Пропускная способность всех трёх алгоритмов равна 0,016 задачи на единицу времени при нулевой цене переключений.',
333:'В модели MQ с seed = 525 завершены все 100 процессов; получены 397 перемещений и 2047 переключений между разными PID. Среднее ожидание составило 17 911,13 мс, использование CPU — 99,9931 %. Учёт обслуживания, ожидания и I/O позволил проверить баланс времени и объяснить отличие фактических долей очередей от календарных бюджетов.'
}
for i,t in repl.items():setp(i,t)
for i in [118,127,136,147,156,165,176,200,209,218,229,238,247,258]:remove(i)

# Native results are transcribed from each screenshot's FIRST_FINISH and SLOT.
observations=[
[['1','9868 / 1300','999000','1000000','1000'],['2','5180 / 6924','1000000','1000000','0'],['3','8088 / 2036','998000','1000000','2000']],
[['1','11756 / 5412','1000000','1000000','0'],['2','6776 / 4720','1000000','1000000','0'],['3','9388 / 9652','1000000','1000000','0']],
[['1','7822 / 7822','1000000','1000000','0'],['2','8067 / 8067','1000000','1000000','0'],['3','8188 / 8188','1000000','1000000','0']],
[['1','8305 / 8307','1000000','1000000','0'],['2','8388 / 8389','1000000','1000000','0'],['3','8496 / 8498','1000000','1000000','0']]
]
for table,rows in zip(d.tables[:4],observations):
 for row,vals in zip(table.rows[1:],rows):
  for c,v in zip(row.cells,vals):c.text=v

# Use a separate editable paragraph for every inline figure and caption.
figures=[]
for i,e in enumerate(body):
 blips=e.xpath('.//a:blip')
 if not blips:continue
 assert len(blips)==1
 b=blips[0];blob=d.part.related_parts[b.get(qn('r:embed'))].blob
 text=''.join(e.xpath('.//w:t/text()'))
 if not text.strip():
  j=i+1
  while not ''.join(body[j].xpath('.//w:t/text()')).strip():j+=1
  cp=Paragraph(body[j],d._body);caption=cp.text
 else:
  caption=text;cp=None
 caption=caption.strip()
 assert caption.startswith('Рисунок'),(i,caption)
 ip=Paragraph(e,d._body);ip.clear()
 if cp is None:
  ce=OxmlElement('w:p');e.addnext(ce);cp=Paragraph(ce,d._body);cp.text=caption
 else:cp.text=caption
 w,h=Image.open(BytesIO(blob)).size
 width=min(165,142*w/h)
 ip.add_run().add_picture(BytesIO(blob),width=Mm(width))
 ip.style='Normal';cp.style='Caption'
 for p in [ip,cp]:
  f=p.paragraph_format;f.first_line_indent=Mm(0);f.left_indent=Mm(0);f.right_indent=Mm(0)
  f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.page_break_before=False;f.keep_together=True;f.keep_with_next=True
  f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=1
 ip.paragraph_format.space_before=Pt(14)
 cp.paragraph_format.space_before=Pt(6);cp.paragraph_format.space_after=Pt(14)
 lead=e.getprevious()
 while lead is not None and lead.tag==qn('w:p') and not ''.join(lead.xpath('.//w:t/text()')).strip():
  prev=lead.getprevious();lead.getparent().remove(lead);lead=prev
 if lead is not None and lead.tag==qn('w:p'):
  lp=Paragraph(lead,d._body);lp.paragraph_format.keep_with_next=True;lp.paragraph_format.keep_together=True
  lp.paragraph_format.space_after=Pt(0)
 after=cp._p.getnext()
 if after is not None and after.tag==qn('w:p'):
  ap=Paragraph(after,d._body);ap.paragraph_format.keep_together=True;ap.paragraph_format.keep_with_next=False
 figures.append({'caption':caption,'sha256':hashlib.sha256(blob).hexdigest(),'width_mm':round(width,2)})

# Restore the self-check appendix referenced by the original contents.
parsed=ast.parse((R/'.work/build_report.py').read_text(encoding='utf8'))
answers=None
for node in parsed.body:
 if isinstance(node,ast.Assign) and any(isinstance(x,ast.Name) and x.id=='answers' for x in node.targets):answers=ast.literal_eval(node.value)
assert answers and len(answers)==15
if not any(p.text.startswith('Приложение Б') and p.style.name=='Heading 1' for p in d.paragraphs):
 d.add_paragraph('Приложение Б Ответы на вопросы для самоконтроля','Heading 1')
 for n,(title,text) in enumerate(answers,1):
  p=d.add_paragraph(f'{n}. {title}. {text}');p.paragraph_format.space_after=Pt(6)

# Respect the supplied formatting rules and preserve the original cover.
for st in ['Normal','Heading 1','Heading 2','Heading 3','Caption','toc 1','toc 2']:
 s=d.styles[st];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.color.rgb=RGBColor(0,0,0)
 for rf in s.element.xpath('.//w:rFonts'):
  for a in list(rf.attrib):
   if 'theme' in a.lower():del rf.attrib[a]
  for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
n=d.styles['Normal'].paragraph_format
n.line_spacing=1;n.first_line_indent=Mm(12.5);n.space_before=Pt(0);n.space_after=Pt(0);n.widow_control=True
for lev in [1,2,3]:
 f=d.styles[f'Heading {lev}'].paragraph_format
 f.space_before=Pt(0 if lev==1 else 18);f.space_after=Pt(18 if lev==1 else 12)
 f.first_line_indent=Mm(0);f.keep_with_next=True;f.keep_together=True;f.page_break_before=lev==1

for p in d.paragraphs[66:]:
 f=p.paragraph_format
 # Remove residual placeholder spacing and forced page breaks from subsections.
 if p.style.name.startswith('Heading'):
  lev=int(p.style.name[-1]);f.page_break_before=lev==1
  f.space_before=Pt(0 if lev==1 else 18);f.space_after=Pt(18 if lev==1 else 12)
  f.keep_with_next=True;f.keep_together=True;f.first_line_indent=Mm(0)
  f.alignment=WD_ALIGN_PARAGRAPH.LEFT
  nxt=p._p.getnext()
  if nxt is not None and nxt.tag==qn('w:p'):
   np=Paragraph(nxt,d._body)
   if np.style.name.startswith('Heading'):
    f.space_after=Pt(12 if lev==1 else 6);np.paragraph_format.space_before=Pt(0)
 else:
  f.page_break_before=False
  if not p._p.xpath('.//w:drawing') and not p.text.startswith(('Рисунок','Таблица')):
   f.space_before=Pt(0);f.space_after=Pt(0);f.first_line_indent=Mm(12.5)
   f.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
   if p.text.strip():f.keep_together=True
 # Manual break characters left behind after users insert screenshots.
 for br in p._p.xpath('.//w:br[@w:type="page"]'):br.getparent().remove(br)
 for r in p.runs:
  r.font.highlight_color=None;r.font.color.rgb=RGBColor(0,0,0);r.font.name='Times New Roman';r.font.size=Pt(14)

# Adjacent heading gaps after the reset above.
for p in d.paragraphs:
 if not p.style.name.startswith('Heading'):continue
 nxt=p._p.getnext()
 if nxt is not None and nxt.tag==qn('w:p'):
  np=Paragraph(nxt,d._body)
  if np.style.name.startswith('Heading'):
   p.paragraph_format.space_after=Pt(12 if p.style.name=='Heading 1' else 6)
   np.paragraph_format.space_before=Pt(0)

widthsets=[[20,49,35,35,26]]*4+[[27,48,48,42],[27,19,31,31,31,26],[24,35,36,34,36],[28,34,49,54],[23,25,27,27,31,32]]
for ti,(t,widths) in enumerate(zip(d.tables,widthsets)):
 t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for col,w in zip(t.columns,widths):col.width=Mm(w)
 for ri,row in enumerate(t.rows):
  row.height=None
  pr=row._tr.get_or_add_trPr()
  if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
  for ci,(c,w) in enumerate(zip(row.cells,widths)):
   c.width=Mm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in c.paragraphs:
    pf=p.paragraph_format;pf.first_line_indent=Mm(0);pf.space_before=Pt(2);pf.space_after=Pt(2);pf.line_spacing=1
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_together=True;pf.keep_with_next=True
    # Use locale-correct decimal separators in editable Russian tables.
    if re.fullmatch(r'\d+\.\d+',p.text):p.text=p.text.replace('.',',')
    if ri==0 and ci==1 and ti<4:p.text='Первый\nPID / TID'
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12 if ri==0 else 14);r.font.highlight_color=None;r.bold=ri==0
 # Keep the complete short table with the first analysis paragraph.
 prev=t._tbl.getprevious()
 if prev is not None and prev.tag==qn('w:p'):
  p=Paragraph(prev,d._body);p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
  p.paragraph_format.first_line_indent=Mm(0);p.paragraph_format.space_before=Pt(14);p.paragraph_format.space_after=Pt(6);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  lead=prev.getprevious()
  if lead is not None and lead.tag==qn('w:p'):
   lp=Paragraph(lead,d._body);lp.paragraph_format.keep_with_next=True;lp.paragraph_format.keep_together=True
  # Each table has an immediate introductory reference, even on a new page.
  if ti==1:
   np=OxmlElement('w:p');prev.addprevious(np)
   lp=Paragraph(np,d._body);lp.text='Результаты потоковых опытов Windows сведены в таблицу 2.2. Во всех вариантах использовался класс процесса Normal, а уровни рабочих потоков задавались отдельно.'
   lp.paragraph_format.keep_with_next=True;lp.paragraph_format.keep_together=True
  elif ti==3:
   np=OxmlElement('w:p');prev.addprevious(np)
   lp=Paragraph(np,d._body);lp.text='Итоги трёх потоковых опытов Linux приведены в таблице 3.2. Первый исполнитель определяется по полю SLOT с учётом соответствующего TID.'
   lp.paragraph_format.keep_with_next=True;lp.paragraph_format.keep_together=True
 nxt=t._tbl.getnext()
 while nxt is not None and nxt.tag==qn('w:p') and not ''.join(nxt.xpath('.//w:t/text()')).strip():
  following=nxt.getnext();nxt.getparent().remove(nxt);nxt=following
 if nxt is not None and nxt.tag==qn('w:p'):
  p=Paragraph(nxt,d._body);p.paragraph_format.space_before=Pt(14);p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=False

# Empty body paragraphs are not used to position content.
for p in list(d.paragraphs)[66:]:
 if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:fldChar|.//w:instrText'):
  p._p.getparent().remove(p._p)

for p in d.paragraphs:
 for hi in p._p.xpath('.//w:highlight'):hi.getparent().remove(hi)
 if re.match(r'[1-8] (Microsoft|Linux man-pages|Arpaci-Dusseau)',p.text):
  p.paragraph_format.first_line_indent=Mm(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(6)
for sect in d.sections:
 sect.page_width=Mm(210);sect.page_height=Mm(297)
 sect.left_margin=Mm(30);sect.right_margin=Mm(15);sect.top_margin=Mm(20);sect.bottom_margin=Mm(20)

alltext='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
assert not re.search(r'\[ЗАПОЛНИТЬ|\[внести\]|МЕСТО ДЛЯ СКРИНШОТА',alltext)
assert len(d.inline_shapes)==37 and len(d.tables)==9
assert len([p for p in d.paragraphs if p.text.startswith('Рисунок ')])==37
d.core_properties.comments='';d.save(OUT)
with ZipFile(BACK) as z:old=sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
with ZipFile(OUT) as z:new=sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
assert old==new,'Screenshot bytes changed'
(W/'changes.json').write_text(json.dumps({'output':str(OUT),'images_unchanged':True,'figures':figures,'observations':observations,'paragraphs':len(d.paragraphs),'tables':len(d.tables)},ensure_ascii=False,indent=2),encoding='utf8')
print(OUT)
print('All 37 original screenshots preserved; 9 tables complete; 15 self-check answers restored.')
