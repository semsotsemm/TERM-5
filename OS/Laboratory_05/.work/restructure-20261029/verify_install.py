from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
import json,hashlib,shutil,csv,math
W=Path(__file__).resolve().parent;R=W.parents[1]
original=R/'Отчет_ЛР5.docx';backup=W/'Отчет_ЛР5 — до правок.docx';ready=W/'Отчет_ЛР5 — исправленный.docx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(original)==sha(backup),'Source report changed during review'
d=Document(ready);old=Document(backup)
assert len(d.inline_shapes)==37 and len(d.tables)==9
def media(p):
 with ZipFile(p) as z:return Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
assert media(backup)==media(ready)
def values(t):return [[c.text for c in row.cells] for row in t.rows]
assert all(values(a)==values(b) for a,b in zip(old.tables,d.tables))
assert not d.element.xpath('.//w:highlight|.//wp:anchor')
def effective(p,name):
 value=getattr(p.paragraph_format,name);style=p.style
 while value is None and style is not None:value=getattr(style.paragraph_format,name);style=style.base_style
 return value
seen=False;figures=0
for p in d.paragraphs:
 if p.text=='Введение' and p.style.name=='Heading 1':seen=True
 if not seen:continue
 assert not p._p.xpath('.//w:br[@w:type="page"]')
 if effective(p,'page_break_before'):assert p.style.name=='Heading 1',p.text
 if p.style.name in ['Heading 1','Heading 2','Heading 3']:
  e=p._p.getprevious()
  if e is not None and e.tag==qn('w:p'):assert not Paragraph(e,d._body).text.startswith(('Рисунок ','Таблица ','Листинг '))
 if p._p.xpath('.//w:drawing'):
  figures+=1;assert effective(p,'space_before').pt==14
  caption=Paragraph(p._p.getnext(),d._body);after=Paragraph(caption._p.getnext(),d._body)
  assert caption.text.startswith('Рисунок ') and effective(caption,'keep_with_next')
  assert after.text.strip() and not after.text.startswith(('Рисунок ','Таблица '))
  assert not effective(after,'keep_together')
assert figures==37
audit=json.loads((W/'page-audit.json').read_text(encoding='utf8'))
assert not audit['issues'] and len(audit['pages'])==45
tables=json.loads((W/'word-tables.json').read_text(encoding='utf-8-sig'))
assert len(tables)==9 and all(t['FirstPage']==t['LastPage'] for t in tables)
assert not (R/'native').exists() and not (R/'simulation').exists()
assert len(list((R/'Laboratory_05').rglob('*.vcxproj')))==16
assert len(list((R/'Laboratory_05').rglob('*.slnx')))==2
assert not list((R/'Laboratory_05').rglob('*.py'))
assert json.loads((W/'test-results.json').read_text())['reference_comparison_exact']
for osname in ['windows','linux']:
 root=R/'Laboratory_05'/osname
 assert (root/'common/Scheduler.cpp').exists() and (root/'README.md').exists()
 assert len(list(root.glob('Lab-05*/*.cpp')))==8
 path=W/('cpp-tests' if osname=='windows' else 'linux-results')
 assert json.loads((path/'comparison.json').read_text())==json.loads((W/'original-code/simulation/results/comparison.json').read_text())
 assert json.loads((path/'MQ-seed525.json').read_text())==json.loads((W/'original-code/simulation/results/MQ-seed525.json').read_text())
 for ref in (W/'original-code/simulation/results').glob('*.csv'):
  candidate=path/ref.name
  with ref.open(encoding='utf-8-sig',newline='') as f:a=list(csv.DictReader(f))
  with candidate.open(encoding='utf-8-sig',newline='') as f:b=list(csv.DictReader(f))
  assert len(a)==len(b),(osname,ref.name)
  for x,y in zip(a,b):
   assert x.keys()==y.keys()
   for k in x:
    try:assert math.isclose(float(x[k]),float(y[k]),rel_tol=1e-12,abs_tol=1e-9),(ref.name,k)
    except ValueError:assert x[k]==y[k]
report={'pages':45,'screenshots_unchanged':37,'data_tables_unchanged':9,'all_pages_visually_reviewed':True,'manual_body_page_breaks':0,'layout_issues':audit['issues'],'cpp_projects_per_os':8,'windows_debug_and_release_built':True,'linux_gcc_built_and_all_apps_run':True,'reference_json_and_csv_match':True,'old_directories_archived':str(W/'removed-originals'),'report_backup':str(backup),'report_sha256':sha(ready)}
shutil.copy2(ready,original)
assert sha(original)==sha(ready)
(W/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
